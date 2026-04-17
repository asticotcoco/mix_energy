from __future__ import annotations

from pathlib import Path
from textwrap import shorten, wrap

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt


ROOT_DIR = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT_DIR / "docs"
REPORT_PATH = DOCS_DIR / "rapport_rncp37827BC01_mix_energy_jean_christophe_charbonnel_soutenance.md"
OUTPUT_PATH = DOCS_DIR / "jcc_mix_energy_soutenance.pptx"
RENDER_DIR = DOCS_DIR / ".pptx_rendered"

TEXT_COLOR = RGBColor(17, 24, 39)
MUTED_TEXT_COLOR = RGBColor(55, 65, 81)
BACKGROUND_COLOR = RGBColor(255, 255, 255)
ACCENT_COLOR = RGBColor(30, 64, 175)

SLIDE_WIDTH_PX = 1600
SLIDE_HEIGHT_PX = 900
ACCENT_HEIGHT_PX = 54
TITLE_TOP_PX = 92
BODY_TOP_PX = 210
LEFT_MARGIN_PX = 96
RIGHT_MARGIN_PX = 96


def read_snippet(relative_path: str, start_line: int, end_line: int) -> str:
    target = ROOT_DIR / relative_path
    lines = target.read_text(encoding="utf-8").splitlines()
    return "\n".join(lines[start_line - 1 : end_line])


def fit_code_block(text: str, max_lines: int = 10, max_width: int = 78) -> str:
    lines = text.splitlines()[:max_lines]
    trimmed: list[str] = []
    for line in lines:
        trimmed.append(shorten(line, width=max_width, placeholder=" ...") if len(line) > max_width else line)
    return "\n".join(trimmed)


def read_report_section(title: str) -> list[str]:
    lines = REPORT_PATH.read_text(encoding="utf-8").splitlines()
    section_lines: list[str] = []
    in_section = False
    target_heading = f"## {title}"

    for line in lines:
        if line.strip() == target_heading:
            in_section = True
            continue
        if in_section and line.startswith("## "):
            break
        if in_section:
            section_lines.append(line)

    return section_lines


def report_section_to_bullets(title: str, limit: int = 4) -> list[str]:
    bullets: list[str] = []
    for line in read_report_section(title):
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("- "):
            bullets.append(stripped[2:])
            continue
        if stripped.startswith("#"):
            continue
        bullets.append(stripped)
        if len(bullets) >= limit:
            break
    return bullets[:limit]


def ensure_render_dir() -> None:
    RENDER_DIR.mkdir(exist_ok=True)
    for image_path in RENDER_DIR.glob("slide_*.png"):
        image_path.unlink()


def rgb_tuple(color: RGBColor) -> tuple[int, int, int]:
    return tuple(color)


def load_font(size: int, bold: bool = False, monospace: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    font_candidates: list[str]
    if monospace:
        font_candidates = [
            "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationMono-Regular.ttf",
        ]
    elif bold:
        font_candidates = [
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf",
        ]
    else:
        font_candidates = [
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
        ]

    for candidate in font_candidates:
        candidate_path = Path(candidate)
        if candidate_path.exists():
            return ImageFont.truetype(str(candidate_path), size=size)
    return ImageFont.load_default()


def wrap_for_pixels(text: str, font, max_width: int) -> list[str]:
    words = text.split()
    if not words:
        return [""]

    lines: list[str] = []
    current = words[0]
    for word in words[1:]:
        trial = f"{current} {word}"
        if font.getlength(trial) <= max_width:
            current = trial
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


def draw_lines(draw: ImageDraw.ImageDraw, lines: list[str], x: int, y: int, font, fill: tuple[int, int, int], line_gap: int) -> int:
    cursor_y = y
    for line in lines:
        draw.text((x, cursor_y), line, font=font, fill=fill)
        bbox = draw.textbbox((x, cursor_y), line, font=font)
        cursor_y = bbox[3] + line_gap
    return cursor_y


def render_slide_image(slide_index: int, title: str, bullets: list[str], code_title: str | None = None, code_text: str | None = None) -> Path:
    image = Image.new("RGB", (SLIDE_WIDTH_PX, SLIDE_HEIGHT_PX), rgb_tuple(BACKGROUND_COLOR))
    draw = ImageDraw.Draw(image)

    draw.rectangle((0, 0, SLIDE_WIDTH_PX, ACCENT_HEIGHT_PX), fill=rgb_tuple(ACCENT_COLOR))

    title_font = load_font(40, bold=True)
    body_font = load_font(26)
    code_title_font = load_font(20, bold=True)
    code_font = load_font(17, monospace=True)

    draw.text((LEFT_MARGIN_PX, TITLE_TOP_PX), title, font=title_font, fill=rgb_tuple(TEXT_COLOR))

    cursor_y = BODY_TOP_PX
    max_width = SLIDE_WIDTH_PX - LEFT_MARGIN_PX - RIGHT_MARGIN_PX
    body_gap = 12

    for bullet in bullets:
        bullet_lines = wrap_for_pixels(bullet, body_font, max_width - 36)
        draw.text((LEFT_MARGIN_PX, cursor_y), "•", font=body_font, fill=rgb_tuple(MUTED_TEXT_COLOR))
        cursor_y = draw_lines(
            draw,
            bullet_lines,
            LEFT_MARGIN_PX + 34,
            cursor_y,
            body_font,
            rgb_tuple(MUTED_TEXT_COLOR),
            body_gap,
        ) + 6

    if code_title and code_text:
        cursor_y += 12
        draw.text((LEFT_MARGIN_PX, cursor_y), f"Code : {code_title}", font=code_title_font, fill=rgb_tuple(TEXT_COLOR))
        cursor_y += 42
        code_box_height = SLIDE_HEIGHT_PX - cursor_y - 60
        draw.rounded_rectangle(
            (LEFT_MARGIN_PX, cursor_y, SLIDE_WIDTH_PX - RIGHT_MARGIN_PX, cursor_y + code_box_height),
            radius=18,
            fill=(244, 247, 250),
        )
        code_lines: list[str] = []
        for raw_line in fit_code_block(code_text, max_lines=11, max_width=88).splitlines():
            wrapped = wrap(raw_line, width=90, replace_whitespace=False, drop_whitespace=False) or [""]
            code_lines.extend(wrapped)
        draw_lines(
            draw,
            code_lines,
            LEFT_MARGIN_PX + 24,
            cursor_y + 22,
            code_font,
            rgb_tuple(TEXT_COLOR),
            6,
        )

    image_path = RENDER_DIR / f"slide_{slide_index:02d}.png"
    image.save(image_path)
    return image_path


def add_base_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    return slide


def add_rendered_slide(prs: Presentation, title: str, bullets: list[str], code_title: str | None = None, code_text: str | None = None) -> None:
    slide = add_base_slide(prs)
    image_path = render_slide_image(len(prs.slides), title, bullets, code_title=code_title, code_text=code_text)
    slide.shapes.add_picture(str(image_path), 0, 0, width=prs.slide_width, height=prs.slide_height)


def add_title_slide(prs: Presentation) -> None:
    add_rendered_slide(
        prs,
        "Mix Energy",
        [
            "Presentation de soutenance basee sur le rapport final personnalise",
            "Jean-Christophe Charbonnel | RNCP37827BC01 | Artefact",
        ],
    )


def add_bullet_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    add_rendered_slide(prs, title, bullets)


def add_code_slide(
    prs: Presentation,
    title: str,
    bullets: list[str],
    code_title: str,
    code_text: str,
) -> None:
    add_rendered_slide(prs, title, bullets, code_title=code_title, code_text=code_text)


def build_deck() -> Presentation:
    ensure_render_dir()
    prs = Presentation()

    add_title_slide(prs)

    add_bullet_slide(
        prs,
        "Objet du document",
        report_section_to_bullets("Objet du document", limit=3),
    )

    add_bullet_slide(
        prs,
        "Fil directeur de soutenance",
        report_section_to_bullets("Note de lecture pour la soutenance", limit=4),
    )

    add_bullet_slide(
        prs,
        "Contexte et objectif",
        [
            "Construire une chaine de valeur complete autour des donnees energetiques francaises.",
            "Centraliser des sources ouvertes heterogenes dans un environnement analytique exploitable.",
            "Calculer le mix electrique, l'empreinte carbone et preparer une prediction de consommation.",
            "Exposer les resultats via une API et un dashboard national et regional.",
        ],
    )

    add_bullet_slide(
        prs,
        "Architecture du projet",
        [
            "Sources : Eco2mix ODRE et Base Carbone ADEME.",
            "Ingestion Python puis depot dans GCS et chargement BigQuery.",
            "Transformations dbt en couches silver et gold.",
            "Orchestration Airflow pour ingestion, chargement et training.",
            "Exposition via FastAPI et restitution via Streamlit.",
        ],
    )

    add_bullet_slide(
        prs,
        "Flux de donnees de bout en bout",
        [
            "1. Recuperation des donnees depuis des APIs ouvertes.",
            "2. Historisation dans un bucket GCS.",
            "3. Chargement des jeux dans BigQuery.",
            "4. Transformation analytique avec dbt.",
            "5. Consommation via API, dashboard et prediction.",
        ],
    )

    add_code_slide(
        prs,
        "Ingestion Python",
        [
            "Le module eco2mix_ingest gere les appels HTTP et les statuts 400, 401, 429 et 500.",
            "Le cas 429 est traite explicitement pour gerer les quotas API.",
            "Les fichiers sont ensuite serialises et envoyes dans le bucket GCP.",
        ],
        "ingest_dbt/src/mix_energy/eco2mix_ingest.py",
        read_snippet("ingest_dbt/src/mix_energy/eco2mix_ingest.py", 14, 44),
    )

    add_code_slide(
        prs,
        "Orchestration Airflow",
        [
            "Le DAG dag_train_model verifie la disponibilite des tables gold avant l'entrainement.",
            "Les taches sont chainees explicitement pour assurer l'ordre des operations.",
            "Cette couche relie data engineering et machine learning.",
        ],
        "airflow/dags/dag_train_model.py",
        read_snippet("airflow/dags/dag_train_model.py", 13, 39),
    )

    add_code_slide(
        prs,
        "Transformations dbt",
        [
            "Le modele reg_tr_predi produit des variables explicatives pour la prediction.",
            "Les fenetres SQL calculent des moyennes glissantes a plusieurs horizons.",
            "Les variables calendaires sont derivees directement en SQL.",
        ],
        "dbt/models/gold/reg_tr_predi.sql",
        read_snippet("dbt/models/gold/reg_tr_predi.sql", 5, 28),
    )

    add_code_slide(
        prs,
        "Exposition FastAPI",
        [
            "L'API parse et valide les filtres avant de lancer les requetes.",
            "La couche web est separee de la couche BigQueryDatasetService.",
            "Les erreurs sont converties en reponses HTTP explicites.",
        ],
        "fastapi/src/mix_energy_api/main.py",
        read_snippet("fastapi/src/mix_energy_api/main.py", 20, 39),
    )

    add_code_slide(
        prs,
        "Service BigQuery securise",
        [
            "Les types BigQuery sont coercis explicitement cote service.",
            "Les colonnes et tables demandees sont verifiees avant execution.",
            "Les query parameters nommes evitent une concatenation SQL naive.",
        ],
        "fastapi/src/mix_energy_api/bigquery_service.py",
        read_snippet("fastapi/src/mix_energy_api/bigquery_service.py", 55, 110),
    )

    add_code_slide(
        prs,
        "Restitution Streamlit",
        [
            "La page d'accueil est legere et ne charge pas les donnees immediatement.",
            "Les pages metier sont separees entre national, regional, historique et temps reel.",
            "Le dashboard exploite les endpoints exposes par FastAPI.",
        ],
        "front-streamlit/dashboard/Accueil_des_dashboards.py",
        read_snippet("front-streamlit/dashboard/Accueil_des_dashboards.py", 10, 40),
    )

    add_code_slide(
        prs,
        "Prediction et ML",
        [
            "Le module train.py charge les donnees, entraine puis evalue le modele.",
            "Le pipeline couvre train, predict et reutilisation des artefacts.",
            "Les metriques sont journalisees via MLflow.",
        ],
        "predict/src/predict/train.py",
        read_snippet("predict/src/predict/train.py", 10, 52),
    )

    add_code_slide(
        prs,
        "Infrastructure Terraform",
        [
            "L'infrastructure GCP est definie en code via Terraform.",
            "Le projet active Compute, BigQuery, Storage, Vertex AI et Artifact Registry.",
            "Des comptes de service dedies sont provisionnes pour les composants du projet.",
        ],
        "iac/main.tf",
        read_snippet("iac/main.tf", 1, 44),
    )

    add_code_slide(
        prs,
        "Tests et qualite logicielle",
        [
            "Les tests unitaires couvrent plusieurs couches du projet.",
            "Le service BigQuery est teste avec des faux objets pour valider le SQL genere.",
            "La robustesse fonctionnelle ne repose pas uniquement sur la demonstration manuelle.",
        ],
        "fastapi/tests/unit/test_bigquery_service.py",
        read_snippet("fastapi/tests/unit/test_bigquery_service.py", 73, 102),
    )

    add_bullet_slide(
        prs,
        "Adequation RNCP",
        [
            "Conception d'une architecture modulaire de bout en bout.",
            "Developpement Python, SQL analytique, orchestration, API et visualisation.",
            "Capacite a industrialiser un projet data sur GCP.",
            "Presence de tests et de mecanismes de securisation des acces aux donnees.",
        ],
    )

    add_bullet_slide(
        prs,
        "Conclusion",
        [
            "Mix Energy constitue une preuve de travail de fond coherent et presentable au jury.",
            "Le projet couvre ingestion, transformation, API, dashboard, prediction et IaC.",
            "La valeur du projet vient de son articulation complete entre besoin metier et implementation technique.",
        ],
    )

    return prs


def main() -> None:
    prs = build_deck()
    prs.save(OUTPUT_PATH)
    print(f"Generated {OUTPUT_PATH}")


if __name__ == "__main__":
    main()